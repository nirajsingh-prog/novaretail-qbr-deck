"""Run:  python -m pytest -q"""
import io
import json
from pathlib import Path

import pytest
from pptx import Presentation

from src import agent_tools as T
from src.analysis import load_and_analyse
from src.narrative import from_json, rule_based, to_json
from src.pipeline import default_paths, initial_state, render
from src.validation import allowed_numbers, ungrounded_numbers, validate_narrative

D = default_paths()


@pytest.fixture(scope="module")
def m():
    return load_and_analyse(D["data"])


class FakeCtx:                       # stands in for google.adk ToolContext
    def __init__(self, state):
        self.state = state


@pytest.fixture
def ctx():
    return FakeCtx(initial_state(D["data"]))


def test_kpis(m):
    assert round(m.kpis["total_revenue_m"]) == 1280
    assert round(m.kpis["revenue_growth_pct"], 1) == 12.3
    assert m.facts["top_growth_region"] == "APAC" and m.facts["bottom_growth_region"] == "Latin America"


def test_json_roundtrip_and_locked_kpis(m):
    j = to_json(rule_based(m), m)
    j["kpis"][0]["value"] = "$9.99B"                 # tampering is ignored
    j["kpis"][0]["label"] = "Q2 Revenue"
    story = from_json(j, m)
    assert story["kpis"][0] == ("$1.28B", "Q2 Revenue")
    assert story[2]["headline"] == j["slide_2"]["headline"]


def test_draft_is_valid(m):
    assert validate_narrative(to_json(rule_based(m), m), allowed_numbers(m)) == []


def test_invented_numbers_are_caught(m):
    A = allowed_numbers(m)
    assert ungrounded_numbers("APAC NPS reached 77", A) == ["77"]
    assert ungrounded_numbers("APAC grew 27% to $330M; group $1.28B", A) == []


def test_structure_rules(m):
    j = to_json(rule_based(m), m)
    j["slide_2"]["bullets"] = j["slide_2"]["bullets"][:2]
    assert any("3-4 bullets" in i for i in validate_narrative(j, allowed_numbers(m)))


def test_agent_tools_edit_state(ctx):
    r = T.update_text_section("slide_2_headline", "APAC Powers Growth; Latin America Needs a Reset", ctx)
    assert r["status"] == "ok"
    assert ctx.state["narrative"]["slide_2"]["headline"].startswith("APAC Powers")
    assert ctx.state["revision_log"][-1]["section"] == "slide_2_headline"


def test_agent_tools_reject_bad_edits(ctx):
    before = json.dumps(ctx.state["narrative"])
    assert T.update_text_section("summary", "Revenue hit $5.4B, up 63%.", ctx)["status"] == "rejected"
    assert T.update_bullets(3, ["only one bullet"], ctx)["status"] == "rejected"
    assert T.update_actions(["Expand"], ["APAC capacity"], ctx)["status"] == "rejected"
    assert T.update_text_section("slide_9_headline", "x", ctx)["status"] == "rejected"
    assert json.dumps(ctx.state["narrative"]) == before          # nothing changed


def test_preferences(ctx):
    T.remember_preference("Keep bullets under 20 words", ctx)
    T.remember_preference("keep bullets under 20 words", ctx)       # de-duplicated
    assert ctx.state["user:preferences"] == ["Keep bullets under 20 words"]
    assert T.get_deck_context(ctx)["user_preferences"] == ["Keep bullets under 20 words"]
    T.forget_preferences(ctx)
    assert ctx.state["user:preferences"] == []


def test_deck_built_from_edited_narrative(m, ctx):
    T.update_text_section("closing", "Back APAC, fix Latin America, and report back at the Q3 2026 QBR.", ctx)
    buf = io.BytesIO()
    log, issues = render(m, ctx.state["narrative"], buf)
    assert not issues and len(log) == 8
    prs = Presentation(io.BytesIO(buf.getvalue()))
    assert len(prs.slides) == 5
    text = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
    assert "fix Latin America" in text and "[FILL" not in text and "[VALUE]" not in text
    assert sum(1 for sl in prs.slides for s in sl.shapes if s.has_chart) == 3
